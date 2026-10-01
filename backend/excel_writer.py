# -*- coding: utf-8 -*-
"""
Excel yazma modülü.
BOM verilerini renk kodlu Excel dosyasına yazan fonksiyonlar.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_FILE

# ============================================================================
# EXCEL STİL TANIMLARI
# ============================================================================

FONT_DATA = Font(name='Calibri', size=10)

BORDER = Border(
    left=Side(style='thin', color='B4C6E7'),
    right=Side(style='thin', color='B4C6E7'),
    top=Side(style='thin', color='B4C6E7'),
    bottom=Side(style='thin', color='B4C6E7'),
)


def write_bom_excel(all_rows):
    """
    Tüm PDF'lerden çıkarılan ve Association Engine'den geçen BOM satırlarını
    teknikerin kolayca inceleyip düzenleyebileceği renk kodlu bir Excel'e yazar.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM & Spool Associations"
    
    # Headers
    headers = [
        "Technical Drawing", "Category", 
        "Assembly (Spool)", "Sub-Assembly (P/L)", 
        "Description", "Item Code", "QTY", "Weight", 
        "Assoc. Status", "Assoc. Method", "Assoc. Reason"
    ]
    
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(name='Calibri', bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = BORDER
        
    ws.row_dimensions[1].height = 25
    
    row_idx = 2
    for r in all_rows:
        trace = r.get("_debug", {})
        status = trace.get("status", "UNKNOWN")
        category = trace.get("category", "")
        
        # Determine background color based on association status
        if status in ["STRONG_DETERMINISTIC", "SPATIAL_HIGH_CONFIDENCE", "VALID"]:
            bg_color = "E2EFDA"  # Light Green
        elif status == "AMBIGUOUS":
            bg_color = "FFF2CC"  # Light Yellow
        elif status == "UNRESOLVED":
            bg_color = "FCE4D6"  # Light Red/Orange
        else:
            bg_color = "FFFFFF"  # White
            
        fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        
        # Prepare row data
        row_data = [
            r.get("technical_drawing", ""),
            category,
            r.get("assembly", ""),
            r.get("sub_assembly", ""),
            r.get("sub_assembly_defination", ""),
            r.get("item_code", ""),
            r.get("qty", ""),
            r.get("unit_weight", ""),
            status,
            trace.get("method", ""),
            trace.get("reason", "")
        ]
        
        for ci, val in enumerate(row_data, 1):
            c = ws.cell(row=row_idx, column=ci, value=val)
            c.font = FONT_DATA
            c.fill = fill
            c.border = BORDER
            c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            
            # Special highlighting for BİLİNMEYEN spools
            if ci == 3 and "BİLİNMEYEN" in str(val).upper():
                c.font = Font(name='Calibri', bold=True, color="FF0000") # Bold Red text
                
        row_idx += 1
        
    # Set column widths
    col_widths = [25, 20, 35, 15, 45, 15, 8, 10, 20, 30, 40]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        
    ws.sheet_properties.tabColor = "2F5496"
    
    # Save
    wb.save(OUTPUT_FILE)
    print(f"\nBOM Excel (Renk Kodlu) kaydedildi: {OUTPUT_FILE}")
